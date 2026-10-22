#!/usr/bin/env python3
"""Build the printable PDF and editable Word index for all fixed Riko page QRs."""

from __future__ import annotations

import json
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
SLOTS_PATH = ROOT / "src/data/bookVideoSlots.json"
QR_DIR = ROOT / "public/qr"
THUMB_DIR = ROOT / "tmp/pdfs/qa-stamped"
PDF_PATH = ROOT / "output/pdf/rikokid-qr-trang-sach.pdf"
DOCX_PATH = ROOT / "output/docx/rikokid-qr-trang-sach.docx"
BG = "#FFFEF0"
BLUE = "#374BA0"
ORANGE = "#F55A43"
PALE = "#F2F5F8"
TEXT = "#3D4858"


def load_slots() -> tuple[dict, list[dict]]:
    manifest = json.loads(SLOTS_PATH.read_text(encoding="utf-8"))
    slots = manifest["slots"]
    for slot in slots:
        if not slot.get("arVideoUrl"):
            raise ValueError(f"No video assigned to {slot['id']}")
        slot["qr_path"] = QR_DIR / f"riko-{slot['id']}.png"
        slot["thumb_path"] = THUMB_DIR / f"ch{slot['id'].split('-')[1][1:]}-{slot['page']:02}.jpg"
        slot["url"] = f"{manifest['publicOrigin']}/watch/{slot['id']}"
        slot["video_file"] = Path(slot["arVideoUrl"]).name
        if not slot["qr_path"].is_file() or not slot["thumb_path"].is_file():
            raise FileNotFoundError(slot["qr_path"] if not slot["qr_path"].is_file() else slot["thumb_path"])
    return manifest, slots


def setup_pdf_fonts() -> tuple[str, str]:
    regular = Path("C:/Windows/Fonts/arial.ttf")
    bold = Path("C:/Windows/Fonts/arialbd.ttf")
    if regular.is_file() and bold.is_file():
        pdfmetrics.registerFont(TTFont("RikoArial", str(regular)))
        pdfmetrics.registerFont(TTFont("RikoArialBold", str(bold)))
        return "RikoArial", "RikoArialBold"
    return "Helvetica", "Helvetica-Bold"


def para(c: canvas.Canvas, text: str, x: float, top: float, width: float, *, size: float = 8, leading: float | None = None, color: str = TEXT, bold: bool = False) -> float:
    style = ParagraphStyle(
        "riko",
        fontName=PDF_BOLD if bold else PDF_REGULAR,
        fontSize=size,
        leading=leading or size * 1.25,
        textColor=colors.HexColor(color),
        alignment=TA_LEFT,
        spaceAfter=0,
        spaceBefore=0,
    )
    p = Paragraph(text, style)
    _, height = p.wrap(width, 200)
    p.drawOn(c, x, top - height)
    return height


def base_page(c: canvas.Canvas, page_no: int, total: int) -> None:
    width, height = A4
    c.setFillColor(colors.HexColor(BG))
    c.rect(0, 0, width, height, fill=1, stroke=0)
    c.setStrokeColor(colors.HexColor("#D7DDE6"))
    c.line(28, 35, width - 28, 35)
    c.setFillColor(colors.HexColor("#697587"))
    c.setFont(PDF_REGULAR, 7)
    c.drawString(28, 22, "RIKOKID  ·  MÃ QR CỐ ĐỊNH THEO TRANG SÁCH")
    c.drawRightString(width - 28, 22, f"{page_no} / {total}")


def draw_cover(c: canvas.Canvas, manifest: dict, slots: list[dict], total: int) -> None:
    base_page(c, 1, total)
    width, height = A4
    c.setFillColor(colors.HexColor(BLUE))
    c.roundRect(28, 626, width - 56, 176, 18, fill=1, stroke=0)
    c.setFillColor(colors.white)
    c.setFont(PDF_BOLD, 9)
    c.drawString(48, 774, "TÀI LIỆU ĐỐI CHIẾU QR")
    c.setFont(PDF_BOLD, 23)
    c.drawString(48, 736, "Mã QR video theo trang sách")
    para(c, escape(manifest["bookTitle"]), 48, 718, width - 96, size=12, color="#FFFFFF")
    para(
        c,
        "Quét mã trong PDF để mở đúng trang video. Mỗi bìa và mỗi trang nội dung có mã riêng.",
        48,
        683,
        width - 96,
        size=9,
        color="#F1F4FF",
    )

    stats = [("3", "file PDF"), (str(len(slots)), "mã QR / trang"), ("54 / 54", "trang có video")]
    x0, gap, card_w = 28, 10, (width - 56 - 20) / 3
    for index, (value, label) in enumerate(stats):
        x = x0 + index * (card_w + gap)
        c.setFillColor(colors.white)
        c.setStrokeColor(colors.HexColor("#D8DFEA"))
        c.roundRect(x, 558, card_w, 52, 12, fill=1, stroke=1)
        c.setFillColor(colors.HexColor(ORANGE if index == 2 else BLUE))
        c.setFont(PDF_BOLD, 17)
        c.drawString(x + 14, 582, value)
        c.setFillColor(colors.HexColor("#697587"))
        c.setFont(PDF_REGULAR, 7)
        c.drawString(x + 14, 568, label)

    para(c, "Danh sách PDF và số trang", 30, 530, 490, size=13, color=BLUE, bold=True)
    y = 500
    for source in manifest["sourceDocuments"]:
        c.setFillColor(colors.HexColor(PALE))
        c.roundRect(28, y - 31, width - 56, 31, 9, fill=1, stroke=0)
        c.setFillColor(colors.HexColor(BLUE))
        c.setFont(PDF_BOLD, 8)
        c.drawString(40, y - 19, f"Chương {source['chapter']}  ·  {source['document']}")
        c.setFillColor(colors.HexColor(TEXT))
        c.setFont(PDF_REGULAR, 7)
        c.drawRightString(width - 40, y - 19, f"{source['pages']} trang = 2 bìa + {source['pages'] - 2} trang ruột")
        y -= 40

    para(c, "Cách dùng", 30, 368, 490, size=13, color=BLUE, bold=True)
    steps = [
        "Mỗi mã QR được in trong PDF đi tới đúng trang tương ứng trên rikokid.me.",
        "Cả hai bìa của mỗi file cũng có QR riêng; tổng cộng có 6 trang bìa.",
        "Clip đã gán cho đủ 54 trang. Các trang bìa và trang kết dùng video Riko; trang tổng kết dùng montage theo nội dung.",
    ]
    y = 340
    for index, step in enumerate(steps, start=1):
        h = para(c, f"{index}. {escape(step)}", 38, y, width - 76, size=8.5, leading=12)
        y -= h + 10

    c.setFillColor(colors.HexColor("#FFF0E8"))
    c.roundRect(28, 66, width - 56, 54, 12, fill=1, stroke=0)
    para(c, "Địa chỉ QR: https://rikokid.me/watch/book-cN-pNN", 42, 105, width - 84, size=8)
    para(c, "Tên file clip hiển thị cạnh từng mã để dễ đối chiếu trước khi in hoặc thay video.", 42, 88, width - 84, size=7.5)
    c.showPage()


def draw_slot_card(c: canvas.Canvas, slot: dict, x: float, y: float, w: float, h: float) -> None:
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor("#D7DDE6"))
    c.roundRect(x, y, w, h, 12, fill=1, stroke=1)
    para(c, escape(slot["label"]), x + 11, y + h - 13, w - 22, size=8.4, color=BLUE, bold=True)
    chapter_no = slot["id"].split("-")[1][1:]
    para(c, f"{escape(slot['document'])} · PDF trang {slot['page']:02} / {next(item['pages'] for item in MANIFEST['sourceDocuments'] if item['chapter'] == int(chapter_no))}", x + 11, y + h - 30, w - 22, size=6.7, color="#697587")

    image_box = (x + 11, y + 42, 112, h - 91)
    ix, iy, iw, ih = image_box
    c.setFillColor(colors.HexColor(PALE))
    c.roundRect(ix, iy, iw, ih, 8, fill=1, stroke=0)
    img = ImageReader(str(slot["thumb_path"]))
    c.drawImage(img, ix + 6, iy + 6, iw - 12, ih - 12, preserveAspectRatio=True, anchor="c", mask="auto")

    qr_x = x + 137
    qr_y = y + h - 145
    c.drawImage(ImageReader(str(slot["qr_path"])), qr_x, qr_y, 91, 91, preserveAspectRatio=True, mask="auto")
    para(c, f"<b>{escape(slot['id'])}</b>", qr_x, qr_y - 8, w - (qr_x - x) - 11, size=6.6, color=BLUE)
    para(c, slot["url"].replace("https://", "https://<br/>"), qr_x, qr_y - 25, w - (qr_x - x) - 12, size=6.4, leading=8, color=TEXT)
    para(c, f"<b>Video:</b> {escape(slot['video_file'])}", qr_x, y + 52, w - (qr_x - x) - 12, size=6.1, leading=8, color="#4A5461")
    para(c, "Quét mã để mở video của trang này.", qr_x, y + 31, w - (qr_x - x) - 12, size=6.1, leading=8, color="#697587")


def build_pdf(manifest: dict, slots: list[dict]) -> None:
    global MANIFEST
    MANIFEST = manifest
    PDF_PATH.parent.mkdir(parents=True, exist_ok=True)
    page_count = 1 + (len(slots) + 3) // 4
    c = canvas.Canvas(str(PDF_PATH), pagesize=A4, pageCompression=1)
    c.setTitle("Mã QR video theo trang sách Riko")
    c.setAuthor("Rikokid")
    draw_cover(c, manifest, slots, page_count)
    width, height = A4
    card_w = (width - 56 - 12) / 2
    card_h = 354
    top = height - 52
    for offset in range(0, len(slots), 4):
        page_no = offset // 4 + 2
        base_page(c, page_no, page_count)
        para(c, "Mã QR · đối chiếu theo trang PDF", 28, height - 22, 400, size=12, color=BLUE, bold=True)
        para(c, f"Các trang {offset + 1:02}–{min(offset + 4, len(slots)):02} / {len(slots)}", width - 200, height - 24, 170, size=7, color="#697587")
        for local_index, slot in enumerate(slots[offset:offset + 4]):
            col = local_index % 2
            row = local_index // 2
            x = 28 + col * (card_w + 12)
            y = top - card_h - row * (card_h + 12)
            draw_slot_card(c, slot, x, y, card_w, card_h)
        c.showPage()
    c.save()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill.replace("#", ""))
    tc_pr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DDE6", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def format_run(run, size=9, color=TEXT, bold=False) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def docx_text(cell, text: str, size=8, color=TEXT, bold=False, after=2):
    p = cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    format_run(run, size, color, bold)
    return p


def add_docx_card(cell, slot: dict) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    set_cell_shading(cell, "FFFFFF")
    set_cell_margins(cell, 100, 110, 80, 110)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    format_run(p.add_run(slot["label"]), 9, BLUE, True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(3)
    p2.paragraph_format.keep_with_next = True
    format_run(p2.add_run(f"{slot['document']} · trang {slot['page']:02}"), 7, "697587")

    inner = cell.add_table(rows=1, cols=2)
    inner.autofit = False
    inner.columns[0].width = Inches(1.62)
    inner.columns[1].width = Inches(1.78)
    left, right = inner.rows[0].cells
    left.width = Inches(1.62)
    right.width = Inches(1.78)
    for subcell in (left, right):
        subcell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        set_cell_margins(subcell, 30, 30, 30, 30)
    left.paragraphs[0].paragraph_format.space_after = Pt(0)
    left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    left.paragraphs[0].add_run().add_picture(str(slot["thumb_path"]), width=Inches(1.48))
    right.paragraphs[0].paragraph_format.space_after = Pt(3)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    right.paragraphs[0].add_run().add_picture(str(slot["qr_path"]), width=Inches(1.16))
    docx_text(right, slot["id"], 7, BLUE, True, 4)
    docx_text(right, slot["url"], 6.2, TEXT, False, 4)
    docx_text(right, f"Video: {slot['video_file']}", 6.5, "4A5461", False, 4)
    docx_text(right, "Quét QR để mở đúng trang video.", 6.2, "697587", False, 0)


def build_docx(manifest: dict, slots: list[dict]) -> None:
    DOCX_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.56)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.4)
    section.right_margin = Inches(0.4)
    section.header_distance = Inches(0.24)
    section.footer_distance = Inches(0.25)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_run(footer.add_run("RIKOKID  ·  MÃ QR CỐ ĐỊNH THEO TRANG SÁCH"), 7, "697587")

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(2)
    format_run(title.add_run("MÃ QR VIDEO THEO TRANG SÁCH"), 22, BLUE, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    format_run(p.add_run(manifest["bookTitle"]), 12, ORANGE, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    format_run(p.add_run("54 mã QR riêng · 54/54 trang có video · gồm đủ 6 trang bìa"), 10, TEXT, True)

    summary = doc.add_table(rows=1, cols=1)
    set_table_borders(summary, "D7DDE6", "7")
    cell = summary.cell(0, 0)
    set_cell_shading(cell, "F2F5F8")
    set_cell_margins(cell, 150, 180, 150, 180)
    cell.text = ""
    for source in manifest["sourceDocuments"]:
        p = cell.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        format_run(p.add_run(f"Chương {source['chapter']} · {source['document']} · {source['pages']} trang (2 bìa)"), 9, BLUE, True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    format_run(p.add_run("Cách dùng"), 12, BLUE, True)
    for step in (
        "Quét mã QR được in trên trang để mở đúng đường dẫn video tại rikokid.me.",
        "Mỗi file PDF có hai trang bìa; cả bìa trước và bìa sau đều có mã riêng.",
        "Tên video bên cạnh QR dùng để kiểm tra trang và clip đã ghép đúng trước khi in.",
    ):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        format_run(p.add_run(step), 9, TEXT)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    format_run(p.add_run("Mã đã được gán video cho toàn bộ 54 trang. Clip mascot dùng cho bìa và trang kết; trang tổng kết dùng montage theo chủ đề."), 8, "4A5461")

    for offset in range(0, len(slots), 4):
        doc.add_page_break()
        page_slots = slots[offset:offset + 4]
        heading = doc.add_paragraph()
        heading.paragraph_format.space_after = Pt(7)
        format_run(heading.add_run(f"Mã QR · đối chiếu theo trang PDF  ({offset + 1:02}–{min(offset + 4, len(slots)):02} / {len(slots)})"), 13, BLUE, True)
        row_count = 2 if len(page_slots) > 2 else 1
        grid = doc.add_table(rows=row_count, cols=2)
        grid.autofit = False
        set_table_borders(grid, "D7DDE6", "7")
        for row in grid.rows:
            if row_count == 2:
                row.height = Inches(4.72)
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
            for cell in row.cells:
                cell.width = Inches(3.68)
        for index, slot in enumerate(page_slots):
            add_docx_card(grid.cell(index // 2, index % 2), slot)
    doc.save(DOCX_PATH)


def main() -> None:
    global PDF_REGULAR, PDF_BOLD
    PDF_REGULAR, PDF_BOLD = setup_pdf_fonts()
    manifest, slots = load_slots()
    build_pdf(manifest, slots)
    build_docx(manifest, slots)
    print(json.dumps({"pdf": str(PDF_PATH), "docx": str(DOCX_PATH), "pages": len(slots), "qrCodes": len(slots)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
